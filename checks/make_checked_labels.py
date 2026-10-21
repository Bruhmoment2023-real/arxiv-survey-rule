"""Build data/labels/labels_checked.csv and data/labels/label_changes_checked.csv from the filled check sheet.

Run:  python -I checks/make_checked_labels.py <repo>     (needs openpyxl)
A first-pass label is replaced only where the check answered N and gave a valid correct label; the change log
lists each change with the checker's reason. data/labels/labels.csv (the first pass) is not touched.
"""
import csv
import sys
from pathlib import Path

from openpyxl import load_workbook

repo = Path(sys.argv[1] if len(sys.argv) > 1 else ".")
wb = load_workbook(repo / "checks" / "09-check-sheet.filled.xlsx", data_only=True)
fix = {}
for name in ("Core 120", "Extra 280 (optional)"):
    for row in wb[name].iter_rows(min_row=2, values_only=True):
        k, yn = row[1], str(row[9] or "").strip().upper()
        corr = str(row[10] or "").strip().upper()
        if k and yn == "N" and corr in set("SPMUN"):
            fix[k] = (corr, " ".join(str(row[11] or "").split()))
lab = list(csv.DictReader(open(repo / "data/labels/labels.csv", encoding="utf-8", newline="")))
changes = []
for r in lab:
    if r["key"] in fix and fix[r["key"]][0] != r["label"]:
        new, why = fix[r["key"]]
        changes.append({"key": r["key"], "first_pass": r["label"], "checked": new, "reason": why})
        r["note"] = (r["note"] + "; " if r["note"] else "") + f"author check: {r['label']}->{new}"
        r["label"], r["labeller"] = new, "AI-assistant-pass1; author check"
with open(repo / "data/labels/labels_checked.csv", "w", encoding="utf-8", newline="") as f:
    w = csv.DictWriter(f, ["key", "label", "labeller", "note"], lineterminator="\n")
    w.writeheader(); w.writerows(lab)
with open(repo / "data/labels/label_changes_checked.csv", "w", encoding="utf-8", newline="") as f:
    w = csv.DictWriter(f, ["key", "first_pass", "checked", "reason"], lineterminator="\n")
    w.writeheader(); w.writerows(changes)
print(f"{len(changes)} labels changed:", ", ".join(f"{c['key']} {c['first_pass']}->{c['checked']}" for c in changes))
