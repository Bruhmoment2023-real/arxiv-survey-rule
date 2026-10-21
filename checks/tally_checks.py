"""Tally Sean's check of the paper 09 first-pass labels.

Run:  D:\\Research-Publishing\\04-sslc-webcasting\\.venv\\Scripts\\python.exe -I tally_checks.py [path\\to\\09-check-sheet.xlsx]
Default path: 09-check-sheet.xlsx next to this script. Needs openpyxl only.

Sean's label for a row: the first-pass label if 'Y'; the 'Correct label' if 'N'; rows marked '?', blank,
or 'N' with no valid correct label are left out (and counted). Positive class = S or P (main), and S, P or M
(sensitivity, as in the paper).

Prints, for the rows checked (core and, separately, everything checked):
  - agreement % (Y among Y+N) with an exact (Clopper-Pearson) 95% interval; same for the binary S/P call
  - disagreements: first-pass label -> Sean's label, and by stratum
  - Cohen's kappa (5 labels, and binary S/P vs not), first pass vs Sean
  - precision and recall of the STRICT rule, by period and pooled, recomputed on the checked rows under the
    first-pass labels and under Sean's labels (same rows, same stratum weights as validate.py), with stratified
    bootstrap 95% intervals, and the pre-registered test: does strict recall fall by > 10 points post vs pre?
This is a non-blind author check of first-pass labels, not an independent blind second labeller.
"""
import math
import random
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

# CS papers on the 57 random days in each stratum (data/labels/strata_sizes.json in the repo)
SIZES = {"strict|pre": 469, "strict|post": 93, "broad_only|pre": 457, "broad_only|post": 203,
         "neither|pre": 13672, "neither|post": 4885}
STRATA = [(s, p) for s in ("strict", "broad_only", "neither") for p in ("pre", "post")]
B = 5000
VALID = set("SPMUN")


def binom_cdf(k, n, p):
    if k < 0:
        return 0.0
    if k >= n:
        return 1.0
    return sum(math.comb(n, i) * p ** i * (1 - p) ** (n - i) for i in range(k + 1))


def cp(k, n):
    """Exact Clopper-Pearson 95% interval by bisection."""
    if n == 0:
        return float("nan"), float("nan")

    def solve(f):
        lo, hi = 0.0, 1.0
        for _ in range(60):
            mid = (lo + hi) / 2
            if f(mid):
                lo = mid
            else:
                hi = mid
        return (lo + hi) / 2
    lower = 0.0 if k == 0 else solve(lambda p: 1 - binom_cdf(k - 1, n, p) < 0.025)
    upper = 1.0 if k == n else solve(lambda p: binom_cdf(k, n, p) > 0.025)
    return lower, upper


def kappa(a, b):
    n = len(a)
    if n == 0:
        return float("nan")
    po = sum(x == y for x, y in zip(a, b)) / n
    ca, cb = Counter(a), Counter(b)
    pe = sum(ca[c] * cb[c] for c in set(ca) | set(cb)) / n / n
    return float("nan") if pe == 1 else (po - pe) / (1 - pe)


def est(rows, periods, lab, pos):
    tp = fp = fn = 0.0
    for p in periods:
        for s in ("strict", "broad_only", "neither"):
            sub = [r for r in rows if r["period"] == p and r["stratum"] == s]
            if not sub:
                return float("nan"), float("nan")
            rate = sum(r[lab] in pos for r in sub) / len(sub)
            n = SIZES[f"{s}|{p}"]
            if s == "strict":
                tp += n * rate
                fp += n * (1 - rate)
            else:
                fn += n * rate
    return (tp / (tp + fp) if tp + fp else float("nan"), tp / (tp + fn) if tp + fn else float("nan"))


def boot(rows, lab, pos, rng):
    strata = {}
    for r in rows:
        strata.setdefault((r["stratum"], r["period"]), []).append(r)
    out = []
    for _ in range(B):
        samp = []
        for v in strata.values():
            samp += rng.choices(v, k=len(v))
        out.append({nm: est(samp, ps, lab, pos) for nm, ps in (("pre", ["pre"]), ("post", ["post"]),
                                                              ("pooled", ["pre", "post"]))})
    return out


def q(xs):
    xs = sorted(x for x in xs if not math.isnan(x))
    if not xs:
        return float("nan"), float("nan")
    return xs[int(0.025 * len(xs))], xs[int(0.975 * len(xs)) - 1]


def load(path):
    wb = load_workbook(path, data_only=True)
    meta = {}
    for row in wb["_key"].iter_rows(min_row=2, values_only=True):
        meta[row[0]] = {"key": row[0], "id": row[1], "period": row[2], "stratum": row[3], "fp": row[5], "set": row[6]}
    rows, skipped = [], Counter()
    for name in ("Core 120", "Extra 280 (optional)"):
        ws = wb[name]
        for row in ws.iter_rows(min_row=2, values_only=True):
            k = row[1]
            if k not in meta:
                continue
            yn = (str(row[9]).strip().upper() if row[9] is not None else "")
            corr = (str(row[10]).strip().upper() if row[10] is not None else "")
            r = dict(meta[k])
            if yn == "Y":
                r["sean"] = r["fp"]
            elif yn == "N" and corr in VALID:
                r["sean"] = corr
                if corr == r["fp"]:
                    skipped["N but correct label equals first pass (treated as Y)"] += 1
            elif yn == "N":
                skipped[f"{name}: N with no valid correct label"] += 1
                continue
            elif yn == "?":
                skipped[f"{name}: ?"] += 1
                continue
            else:
                skipped[f"{name}: blank"] += 1
                continue
            r["note"] = row[11]
            rows.append(r)
    return rows, skipped


def report(rows, title, rng):
    print(f"\n=== {title}: {len(rows)} rows with a usable answer ===")
    if not rows:
        return
    n = len(rows)
    k = sum(r["sean"] == r["fp"] for r in rows)
    lo, hi = cp(k, n)
    print(f"Agreement, 5 labels: {k}/{n} = {k / n:.1%} (exact 95% {lo:.1%}-{hi:.1%})")
    for posname, pos in (("S/P", "SP"), ("S/P/M", "SPM")):
        kb = sum((r["sean"] in pos) == (r["fp"] in pos) for r in rows)
        lo, hi = cp(kb, n)
        print(f"Agreement, binary {posname} vs not: {kb}/{n} = {kb / n:.1%} (exact 95% {lo:.1%}-{hi:.1%})")
    print(f"Cohen's kappa, 5 labels: {kappa([r['fp'] for r in rows], [r['sean'] for r in rows]):.3f}")
    kb = kappa([r["fp"] in "SP" for r in rows], [r["sean"] in "SP" for r in rows])
    kbs = [kappa(*zip(*[(x["fp"] in "SP", x["sean"] in "SP") for x in rng.choices(rows, k=n)])) for _ in range(2000)]
    print(f"Cohen's kappa, binary S/P vs not: {kb:.3f} (bootstrap 95% {q(kbs)[0]:.3f}-{q(kbs)[1]:.3f})")
    print("Labels, first pass:", dict(Counter(r["fp"] for r in rows)), " Sean:", dict(Counter(r["sean"] for r in rows)))
    dis = [r for r in rows if r["sean"] != r["fp"]]
    print(f"Disagreements ({len(dis)}), first pass -> Sean:",
          dict(Counter(f"{r['fp']}->{r['sean']}" for r in dis)) or "none")
    print("Disagreements by stratum:", dict(Counter(f"{r['stratum']}|{r['period']}" for r in dis)) or "none")
    for r in dis:
        print(f"   {r['key']} {r['id']} {r['stratum']}|{r['period']} {r['fp']}->{r['sean']}  {r['note'] or ''}")
    have = Counter((r["stratum"], r["period"]) for r in rows)
    print("Rows per stratum:", {f"{s}|{p}": have[(s, p)] for s, p in STRATA})
    if any(have[sp] == 0 for sp in STRATA):
        print("Precision/recall skipped: at least one of the 6 strata has no checked rows.")
        return
    print("Note: recall leans on the 'neither' strata (97% of the stratum weight), where positives are rare; on the "
          "core alone it is noisy. Compare first pass vs Sean on the SAME rows; the paper's recall uses all 400.")
    for posname, pos in (("main S/P", "SP"), ("sensitivity S/P/M", "SPM")):
        print(f"\nStrict rule, positive = {posname} (stratum-weighted; bootstrap {B}, 95%)")
        print(f"{'labels':<11}{'period':<8}{'precision':<26}{'recall':<26}")
        for lab, lname in (("fp", "first pass"), ("sean", "Sean")):
            draws = boot(rows, lab, pos, rng)
            for nm, ps in (("pre", ["pre"]), ("post", ["post"]), ("pooled", ["pre", "post"])):
                pr, rc = est(rows, ps, lab, pos)
                cpr = q([d[nm][0] for d in draws])
                crc = q([d[nm][1] for d in draws])
                print(f"{lname:<11}{nm:<8}{pr:.3f} ({cpr[0]:.3f}-{cpr[1]:.3f})      "
                      f"{rc:.3f} ({crc[0]:.3f}-{crc[1]:.3f})")
            d = est(rows, ["post"], lab, pos)[1] - est(rows, ["pre"], lab, pos)[1]
            cd = q([x["post"][1] - x["pre"][1] for x in draws])
            flag = "YES - pre-registered rule triggered" if d < -0.10 else "no"
            print(f"{lname:<11}recall post minus pre: {d:+.3f} ({cd[0]:+.3f} to {cd[1]:+.3f}); fall > 10 points? {flag}")


def main():
    path = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parent / "09-check-sheet.xlsx"
    rows, skipped = load(path)
    print(f"Sheet: {path}")
    print("Not used:", dict(skipped) or "none")
    rng = random.Random(20261021)
    report([r for r in rows if r["set"] == "core"], "Core 120 (30% of each stratum)", rng)
    if any(r["set"] == "extra" for r in rows):
        extra_full = sum(r["set"] == "extra" for r in rows) == 280
        report(rows, "All checked rows" + ("" if extra_full else
               " (WARNING: Extra 280 only partly done; strata are then sampled unevenly, which the stratum weights"
               " still handle, but any row-picking by Sean is not random)"), rng)
    print("\nDescribe this as a non-blind author check of the first-pass labels, not a blind second labeller.")


if __name__ == "__main__":
    main()
